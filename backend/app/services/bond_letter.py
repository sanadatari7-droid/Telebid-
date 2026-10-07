"""Bid bond request letter, filled into the company's Word template
(app/templates/bid_bond_request.docx — made from the company's own Bid_Bond_T.docx).

The letter's fixed wording (title, To, From, opening, closing, notes) is edited once in
Company Settings → Bid Bond Approval Cycle; each bond fills the RFP-info table, and the
sign-off table is rebuilt for the approval levels with each approver's name and date.
"""
import io
from copy import deepcopy
from datetime import date
from pathlib import Path

import docx
from docx.oxml.ns import qn

TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "bid_bond_request.docx"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

LETTER_DEFAULTS = {
    "letter_title": "REQUEST for Bid BONDS",
    "letter_to": "",
    "letter_from": "",
    "letter_intro": "We are in the process of delivering a solution and a Bid Bond (BB) is required to be issued as per the details below:",
    "letter_requester_title": "Manager Bids",
    "letter_closing": "Best Regards,",
    "letter_notes": "",
}

_PCT_WORDS = {1: "One", 2: "Two", 3: "Three", 4: "Four", 5: "Five", 10: "Ten"}


def _set_text(paragraph, text: str):
    """Replace a paragraph's text but keep its first run's formatting; '\\n' becomes a line break."""
    runs = paragraph.runs
    run = runs[0] if runs else paragraph.add_run()
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    lines = (text or "").split("\n")
    run.text = lines[0]
    for line in lines[1:]:
        run.add_break()
        run.add_text(line)


def _set_cell(cell, text: str):
    paragraphs = cell.paragraphs
    for p in paragraphs[1:]:
        p._p.getparent().remove(p._p)
    _set_text(paragraphs[0], text)


def _para(document, starts_with: str):
    return next(p for p in document.paragraphs if p.text.strip().startswith(starts_with))


def _date(d) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def _money(value, decimals: int) -> str:
    return f"{float(value):,.{int(decimals)}f}" if value is not None else ""


def _value_and_pct(bond: dict, currency: str, decimals: int) -> str:
    pct = bond.get("lg_percentage")
    if pct is None:
        return f"{currency} {_money(bond.get('bond_amount'), decimals)}".strip()
    p = float(pct)
    words = _PCT_WORDS.get(int(p)) if p == int(p) else None
    pct_text = f"{words} Percent ({p:g}%)" if words else f"{p:g}%"
    return (f"{currency} {_money(bond.get('bond_amount'), decimals)} — "
            f"{pct_text} of {currency} {_money(bond.get('lg_base_value'), decimals)}").strip()


def _resize_signoff(table, columns: int):
    """The template's sign-off table has Requester + 2 approval columns; make it Requester + N."""
    tbl = table._tbl
    grid = tbl.tblGrid
    total = sum(int(g.get(qn("w:w"))) for g in grid.gridCol_lst)
    while len(grid.gridCol_lst) < columns:
        grid.gridCol_lst[-1].addnext(deepcopy(grid.gridCol_lst[-1]))
    while len(grid.gridCol_lst) > columns:
        grid.remove(grid.gridCol_lst[-1])
    width = total // columns
    for g in grid.gridCol_lst:
        g.set(qn("w:w"), str(width))
    rows = tbl.tr_lst
    for tr in rows[1:]:
        while len(tr.tc_lst) < columns:
            tr.tc_lst[-1].addnext(deepcopy(tr.tc_lst[-1]))
        while len(tr.tc_lst) > columns:
            tr.remove(tr.tc_lst[-1])
    header_approvals = rows[0].tc_lst[1]
    tc_pr = header_approvals.get_or_add_tcPr()
    span = tc_pr.find(qn("w:gridSpan"))
    if span is None:
        span = tc_pr.makeelement(qn("w:gridSpan"), {})
        tc_pr.insert(0, span)
    span.set(qn("w:val"), str(columns - 1))
    for tr in rows:
        for i, tc in enumerate(tr.tc_lst):
            tcw = tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcw is not None:
                tcw.set(qn("w:w"), str(width * (columns - 1) if (tr is rows[0] and i == 1) else width))
                tcw.set(qn("w:type"), "dxa")


def build_request_letter(bond: dict, letter: dict, levels: list, currency: str, decimals: int,
                         letter_date: date) -> bytes:
    """levels: [(level_title, approver_name or None, approved_at or None), ...] in order."""
    text = {**LETTER_DEFAULTS, **{k: v for k, v in (letter or {}).items() if v is not None}}
    d = docx.Document(str(TEMPLATE))

    _set_text(_para(d, "REQUEST"), text["letter_title"])
    _set_text(_para(d, "To:"), f"To: {text['letter_to']}".rstrip())
    _set_text(_para(d, "From:"), f"From: {text['letter_from']}".rstrip())
    _set_text(_para(d, "We are in the process"), text["letter_intro"])
    _set_text(_para(d, "Best Regards"), text["letter_closing"])
    notes = "\n".join(n for n in (text["letter_notes"], bond.get("notes")) if n)
    _set_text(_para(d, "Notes:"), f"Notes: {notes}".rstrip())
    for p in d.sections[0].header.paragraphs:
        if p.text.strip().startswith("Date:"):
            _set_text(p, f"Date: {letter_date:%B} {letter_date.day}, {letter_date.year}")

    info, signoff = d.tables
    values = [
        bond.get("bid_ref") or "",
        bond.get("bid_subject") or "",
        bond.get("beneficiary") or "",
        bond.get("beneficiary_address") or "",
        _value_and_pct(bond, currency or "", decimals),
        _date(bond.get("expiry_date")),
        bond.get("language") or "",
        _date(bond.get("submission_date")),
    ]
    for row, value in zip(info.rows[1:], values):
        _set_cell(row.cells[1], value)

    _resize_signoff(signoff, 1 + len(levels))
    roles, names = signoff.rows[1].cells, signoff.rows[2].cells
    _set_cell(roles[0], text["letter_requester_title"])
    _set_cell(names[0], bond.get("requester_name") or "")
    for i, (title, name, at) in enumerate(levels, 1):
        _set_cell(roles[i], title)
        _set_cell(names[i], f"{name}\nApproved {_date(at)}" if name else "Pending")

    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def letter_filename(bond: dict) -> str:
    # ASCII only: it goes into an HTTP header and an email attachment name.
    ref = "".join(ch if (ch.isascii() and ch.isalnum()) or ch in "-_ " else "-" for ch in (bond.get("bid_ref") or str(bond.get("bond_id") or "")))
    return f"Bid Bond Request - {ref.strip() or 'draft'}.docx"
